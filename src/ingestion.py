import os
import re
from typing import List, Dict, Any
from dataclasses import dataclass
import hashlib

@dataclass
class Document:
    document_id: str
    source: str
    content: str
    metadata: Dict[str, Any]

def generate_document_id(source: str) -> str:
    return hashlib.md5(source.encode('utf-8')).hexdigest()[:12]

class DocumentIngestionPipeline:
    def __init__(self, data_dir: str):
        self.data_dir = data_dir

    def load_documents(self) -> List[Document]:
        documents = []
        for root, _, files in os.walk(self.data_dir):
            for file in files:
                if file.startswith('.'):
                    continue
                file_path = os.path.join(root, file)
                relative_path = os.path.relpath(file_path, self.data_dir)
                
                # Extract metadata from path (e.g. "ZX Bank/pdf/filename.pdf")
                parts = relative_path.split(os.sep)
                organization = parts[0] if len(parts) > 0 else "Unknown"
                
                ext = file.lower().split('.')[-1]
                content = ""
                
                try:
                    if ext == "md" or ext == "txt":
                        with open(file_path, "r", encoding="utf-8") as f:
                            content = f.read()
                    elif ext == "pdf":
                        # pyrefly: ignore [missing-import]
                        import fitz # PyMuPDF
                        doc = fitz.open(file_path)
                        text_pages = []
                        for page in doc:
                            text_pages.append(page.get_text())
                        content = "\n\n".join(text_pages)
                    elif ext == "docx":
                        # pyrefly: ignore [missing-import]
                        import docx
                        doc = docx.Document(file_path)
                        content = "\n".join([para.text for para in doc.paragraphs])
                    elif ext == "pptx":
                        # pyrefly: ignore [missing-import]
                        import pptx
                        prs = pptx.Presentation(file_path)
                        text_runs = []
                        for slide in prs.slides:
                            for shape in slide.shapes:
                                if hasattr(shape, "text"):
                                    text_runs.append(shape.text)
                        content = "\n\n".join(text_runs)
                    elif ext == "html" or ext == "htm":
                        from bs4 import BeautifulSoup
                        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                            soup = BeautifulSoup(f.read(), "lxml")
                            content = soup.get_text(separator="\n")
                    else:
                        continue # Skip unknown formats
                except Exception as e:
                    print(f"Error parsing {file_path}: {e}")
                    continue
                    
                cleaned_content = self.clean_content(content)
                if not cleaned_content:
                    continue
                
                doc_id = generate_document_id(relative_path)
                
                metadata = {
                    "document_id": doc_id,
                    "source": relative_path, # keep the original relative path as source
                    "source_file": file,
                    "organization": organization,
                    "file_type": ext
                }
                
                doc = Document(
                    document_id=doc_id,
                    source=relative_path,
                    content=cleaned_content,
                    metadata=metadata
                )
                documents.append(doc)
        return documents

    def clean_content(self, content: str) -> str:
        # Remove markdown comments (if any)
        content = re.sub(r'<!--.*?-->', '', content, flags=re.DOTALL)
        # Collapse multiple newlines
        content = re.sub(r'\n{3,}', '\n\n', content)
        return content.strip()
